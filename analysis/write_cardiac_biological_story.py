"""Build the biology-led Word manuscript with integrated current plates."""
from pathlib import Path
import re,json,datetime
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
import write_cardiac_story as style

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'results/cardiac_biological_synthesis'
SOURCE=ROOT/'analysis/cardiac_biological_story.md'
FIGURES=[
('Figure_1_CM_identity_and_maturation','CM identity and maturation vary independently within early hearts',
 'a,b, Named ventricular-associated cardiomyocyte states at P0 and P7. c,d, Percentiles of the directly measured 21-gene maturation index within each section. e,f, Score distributions within each identity. Labels retain the released reference identities. These are relative transcriptional scores, not percentages of physiological maturity. Focal selection, coordinates and scores are unchanged from the preceding subtype analysis. Scale bars, 500 µm.'),
('Figure_2_fibroblast_development_and_assignment','Fibroblast programs change as the postnatal heart develops',
 'a–d, Whole-cell/boundary-core multi-gene fibroblast candidates on the four native sections; gray shows resolved broad-lineage background. Nuclear support is optional; supported conflicting nuclear evidence vetoes assignment. Missing individual anchor detection alone does not reject these candidates. Unresolved profiles remain in the atlas. e, Direct native Postn detection among eligible broad fibroblasts; residual detection is not proof of a distinct late Postn family or ambient-free expression. f, Assignment coverage under three evidence layers. Counts are operational coverage rather than unbiased subtype abundance. Scale bars, 500 µm.'),
('Figure_3_fibroblast_matrix_specialization','Later fibroblast states differ in matrix-associated programs',
 'The earlier integrated fibroblast plate is retained without changing its data: confidence-supported family maps, full-RNA library-by-family matrix programs, named P21 RNA substates and selected marker-withheld cross-stage signature similarities. These spatial maps retain the earlier stringent family version; the broader candidate maps are in Figure 2. Higher basement-matrix RNA in late Col8a1 versus Gfpt2 families motivates a matrix-organization hypothesis. Selected early/late signature similarities nominate tracing comparisons, rather than measured family conversion or shared ancestry. Detailed source values accompany the previous release.'),
('Figure_4_vascular_growth_and_CM_specialization','Vascular growth and cardiomyocyte specialization have distinct developmental relationships',
 'a, Synthesis of the regional vascular sequence demonstrated by Lu et al. (2021); not newly observed perfusion. b,c, Within-region component associations under legacy, nuclear-optional and multi-gene fibroblast source layers, with 95% within-section block intervals. EC sources are unchanged. The P7 Postn/metabolic association weakens under broader FB definitions; the P14 Ankrd1/endocardial sarcomere association persists. The 30-µm estimates use geometry-specific cohorts. d, Earlier matched capillary-apposition component differences in a broader CM cohort; these are a separate analysis. e, Distinctions needed to interpret growth, access and function. No heart-level replication or contact mechanism is established.'),
('Figure_5_CM_identity_conditioned_transport','Developmental correspondence can preserve some regional ordering',
 'Frozen transport plate: a, Identity-conditioned comparison design. b, Original stringent FB/EC descriptors in 96 reserved observed neighborhoods. c,d, Saved Wormhole approximation versus exact distances. e, Within-stage same- versus different-CM-identity neighborhood distances. f, Held-out relative maturation-rank correspondence by developmental interval and endocardial-distance stratum. New fibroblast candidates do not retrain or alter the distances or plans. Fine matches are variable. Uniform correspondence mass is not ancestry, cell conversion, expansion or developmental time; selected rank alignment remains exploratory.'),
('Figure_6_regional_developmental_model','A regional developmental model leads to discriminating experiments',
 'a, Proposed synthesis of changing early growth, matrix and vascular contexts followed by later specialization. Arrows denote developmental windows, not measured cell conversion or a proven causal chain. b, Supply, matrix/mechanics and cellular integration provide distinct, potentially cooperating explanations for CM development. c, Independent intermediate-stage hearts, matched functional perturbations and lineage tracing distinguish these explanations. Confirm actual vascular perfusion and starting CM identity; measure respiration, calcium and force alongside RNA. The model integrates our secondary analysis with published vascular biology and remains a testable interpretation.')]
SUPPLEMENT=('Figure_S1_fibroblast_assignment_audit','What the fibroblast reassessment does and does not resolve',
 'a, Eligible Gfpt2 source counts under public, whole-cell, boundary-supported and legacy nuclear-required evidence. b, Whole-cell family proposals among eligible P21 public Gfpt2 cells, before nuclear corroboration. Disagreement at this step cannot be explained solely by sparse nuclei. c, Corrected snRNA Postn detection across library-by-family pseudobulks; each point is a family/library summary, not an independent heart-level test. d, Interpretation of the revised evidence layers. Marker and reference agreement do not establish anatomical ground truth.')

def blocks(text):
    return [x.strip() for x in re.split(r'\n\s*\n',text.strip()) if x.strip()]
def add_blocks(doc,text,compact=False):
    in_methods=False;in_discussion=False
    for b in blocks(text):
        if b.startswith('### '):doc.add_heading(b[4:],2)
        elif b.startswith('## '):
            title=b[3:]
            in_methods=title=='Methods'
            in_discussion=title=='Discussion'
            p=doc.add_heading(title,1)
            if title in ['Introduction','Discussion','Methods','References']:p.paragraph_format.page_break_before=True
        elif b.startswith('# '):continue
        else:
            p=doc.add_paragraph(b.replace('\n',' '))
            if in_methods or compact:
                p.paragraph_format.space_after=Pt(6)
                for r in p.runs:r.font.size=Pt(10.5)
            elif in_discussion:
                p.paragraph_format.line_spacing=1.03;p.paragraph_format.space_after=Pt(6)

def plate(doc,entry,number):
    stem,title,legend=entry;p=doc.add_heading(f'Figure {number}. {title}',2);p.paragraph_format.page_break_before=True
    p=OUT/'figures'/(stem+'.png')
    with Image.open(p) as im:w,h=im.size
    width=min(7.0,7.15*w/h)
    paragraph=doc.add_paragraph();paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after=Pt(4)
    paragraph.add_run().add_picture(str(p),width=Inches(width))
    p=doc.add_paragraph(legend);p.paragraph_format.space_after=Pt(0)
    for r in p.runs:r.font.size=Pt(9)

def main():
    text=SOURCE.read_text();title=text.splitlines()[0][2:];doc=Document()
    props=doc.core_properties;props.title=title;props.author='Heart atlas analysis';props.subject='Biological synthesis of cardiac muscle, matrix and vascular development'
    props.created=datetime.datetime(2026,10,26,12,tzinfo=datetime.timezone.utc);props.modified=props.created
    s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=Inches(.65);s.bottom_margin=Inches(.65);s.left_margin=Inches(.7);s.right_margin=Inches(.7)
    s.header_distance=Inches(.25);s.footer_distance=Inches(.25);s.different_first_page_header_footer=True
    n=doc.styles['Normal'];n.font.name='Cambria';n.font.size=Pt(11);n.paragraph_format.line_spacing=1.08;n.paragraph_format.space_after=Pt(7)
    for name in ['Title','Heading 1','Heading 2']:
        doc.styles[name].font.name='Arial';doc.styles[name].font.color.rgb=RGBColor.from_string('203C54')
    doc.styles['Heading 1'].font.size=Pt(15);doc.styles['Heading 2'].font.size=Pt(11.5)
    header=s.header.paragraphs[0];header.text='POSTNATAL HEART DEVELOPMENT | MUSCLE, MATRIX AND VASCULAR COORDINATION'
    for r in header.runs:r.font.name='Arial';r.font.size=Pt(8)
    p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
    doc.add_paragraph('BIOLOGICAL SYNTHESIS • REVISED MANUSCRIPT','Subtitle')
    doc.add_paragraph(title,'Title')
    doc.add_paragraph('An integrated secondary analysis of P0–P21 snRNA-seq and Xenium data')
    doc.add_heading('Central message',2)
    central='Postnatal heart development involves regionally coordinated changes in muscle, matrix and vasculature. Cardiomyocyte identities contain different maturation states, and their changing environments may support different aspects of growth and specialization at different times.'
    doc.add_paragraph(central)
    doc.add_paragraph('This revision emphasizes the developmental biology, integrates the perinatal vascular study, and reassesses fibroblast assignments without compulsory nuclear confirmation. Six integrated main figures and an assignment-audit supplement are included. One spatial section per age limits developmental and causal inference.')
    doc.add_paragraph('Prepared 26 October 2026 • Original atlas and earlier releases retained for traceability.')
    doc.add_page_break()
    body,methods=text.split('## Methods',1)
    add_blocks(doc,body)
    for i,entry in enumerate(FIGURES,1):plate(doc,entry,i)
    add_blocks(doc,'## Methods'+methods)
    doc.add_page_break();doc.add_heading('Evidence and experimental priorities',1)
    style.table(doc,['Biological inference','Evidence status','Next discriminating measurement'],[
        ['CM identity and maturation vary independently within early sections','Measured RNA heterogeneity; function not measured','Subtype-matched respiration, sarcomeres, calcium and force'],
        ['Fibroblast programs specialize with development','Reviewed RNA and candidate spatial placement','Matrix protein, organization and mechanical measurements'],
        ['Perinatal coronary expansion is regionally concentrated','Published lineage tracing and vessel measurements [2]','Matched inner/outer anatomy and perfusion in independent hearts'],
        ['P7 cycling-CM/Postn context accompanies lower metabolic RNA','Direction retained but weaker after reassignment; exploratory','Matched fibroblast/matrix environment and starting CM state'],
        ['P14 Ankrd1/endocardial context accompanies sarcomere RNA','Persists in broader-FB source sensitivity; observational','Anatomy/load control and EC-type-by-CM-type functional comparison'],
        ['Some early/late environments or RNA states resemble one another','Transport/signature candidate; no ancestry established','Intermediate stages and validated lineage tracing']])
    plate(doc,SUPPLEMENT,'S1')
    report=OUT/'Heart_Atlas_segmentation_assessment.md'
    if report.exists():
        doc.add_page_break();doc.add_heading('Supplementary assessment: segmentation in the Heart Atlas',1);add_blocks(doc,report.read_text(),compact=True)
    path=OUT/'Postnatal_heart_biological_synthesis.docx';doc.save(path)
    md=text+'\n\n## Integrated figures\n\n'
    for i,(stem,t,leg) in enumerate(FIGURES,1):md+=f'### Figure {i}. {t}\n\n![Figure {i}](figures/{stem}.png)\n\n{leg}\n\n'
    (OUT/'Manuscript_text.md').write_text(md.rstrip()+'\n')
    (OUT/'manuscript_structure.json').write_text(json.dumps({'title':title,'central_message':central,'main_figures':[x[0] for x in FIGURES],
        'supplementary_figures':[SUPPLEMENT[0]],'main_text_words':len(body.split()),'methods_references_words':len(methods.split()),
        'new_biological_replicates':0,'nuclear_optional_FB_sensitivity':True,'old_atlas_or_transport_overwritten':False},indent=2)+'\n')
    print('BIOLOGICAL_MANUSCRIPT_CREATED',path,flush=True)
if __name__=='__main__':main()
