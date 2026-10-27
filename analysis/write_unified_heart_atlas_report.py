"""Write one integrated report with current conclusions and consolidated tables."""
from pathlib import Path
import json,os,datetime,copy
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from write_cardiac_biological_story import add_blocks
import write_cardiac_story as style
from write_fibroblast_CM_refit_report import LEGENDS as REFIT_LEGENDS
from write_clear_heart_atlas_report import SEGMENTATION_NOTE
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'results/heart_atlas_unified_report';SOURCE=ROOT/'analysis/heart_atlas_unified_report.md'
BIO='results/cardiac_biological_synthesis/figures/';REFIT='results/cardiac_fibroblast_CM_refit/figures/';FB='results/fibroblast_multimethod_identity/figures/';HERE='results/heart_atlas_unified_report/figures/'
def refit(old,new):
 number,stem,title,legend=next(r for r in REFIT_LEGENDS if r[0]==old);return (new,REFIT+stem,title,legend)
FIGURES=[
('1',HERE+'Figure_1_cardiac_lineages_and_endothelial_subtypes','Cardiac identities and endothelial subtypes define different tissue settings',
 'a–d, Broad cardiac identities at P0, P7, P14 and P21 on the original coordinates. Dense populations may be subsampled for display; calculations retain the full cohorts. e–h, Supported capillary-, endocardial-, arterial-, venous- and lymphatic-like endothelial sources. Unassigned eligible ECs and background tissue remain gray. Sparse calls do not establish absence of a vessel type. These maps retain the reviewed broad/EC definitions and do not newly classify cells. Scale bars, 500 µm.'),
('2',BIO+'Figure_1_CM_identity_and_maturation','CMs of the same subtype can differ in developmental state',
 'a,b, Ventricular-associated CM identities at P0 and P7. c,d, Relative position in each section’s measured 21-gene score distribution. e,f, Scores within the named CM groups. Higher scores indicate stronger expression of the measured panel, rather than a direct measurement of contraction or energy use. Subtype and developmental state are complementary descriptions. The underlying labels, scores and coordinates are unchanged. Scale bars, 500 µm.'),
('3',HERE+'Figure_3_RNA_matrix_and_updated_spatial_context','Fibroblast RNA states differ in matrix-related features',
 'a, Full-RNA matrix programs across age and family, averaged over the two libraries and shown relative to the displayed groups. b, P21 marker patterns in named Col8a1/Svep1, Col8a1/Hmcn2, Lsamp/Slit3 and Lsamp/Col4a5 groups; dot size gives detection and color gives relative expression. c, Latest broader-candidate P14 Col8a1/Gfpt2 technical-matched comparison, replacing the earlier stringent estimate. Arterial proximity and smooth-muscle-neighbor share are separate measures. d, Selected early/late RNA similarities with interpreted markers withheld. Dashed arrows nominate comparisons for tracing, not observed cell conversion.'),
refit('1','4'),
('5',FB+'Figure_2_P21_identity_and_program_maps','A Gfpt2-associated pattern extends beyond the narrow candidate subtype',
 'a, The paper’s fibroblast labels within the eligible P21 cohort. b, Candidate identities from agreement between classifiers. c, A smaller subset supported in the cell core and after naming genes are hidden. d, The measured Gfpt2/Pi16/Adgrd1 pattern relative to matched controls. e,f, Gfpt2- and Col8a1-associated patterns can overlap across all eligible fibroblasts and within the public Gfpt2 group. Shared expression does not establish one exclusive subtype or lineage transition. Kirrel3 candidates are not assigned valve anatomy. Scale bars, 500 µm.'),
('6',HERE+'Figure_6_coronary_growth_and_CM_components','Vascular growth and proximity relate to different parts of CM development',
 'a, The regional coronary sequence described by Lu et al. (2021); perfusion was not newly measured here. b, Earlier nuclear-QC EC/CM cell-cycle associations within 50 µm are positive at P0/P7. c, Earlier locally matched capillary-adjacency comparisons show lower CM cycling-associated RNA at P7 and higher metabolic-regulatory RNA at P14/P21. These estimates use their original cohorts and uncertainty definitions and were not rerun with the fibroblast refit. Intervals are within-section spatial-block intervals, not independent-heart replication. Vessel proximity does not establish oxygen delivery or a contact mechanism.'),
refit('4','7'),refit('5','8'),refit('6','9')]
SUPPLEMENT=[refit('2','S1'),refit('3','S2'),
('S3',HERE+'Figure_S3_transport_calibration_and_sampling','Cross-age neighborhood matches have useful but limited resolution',
 'a, Agreement between frozen Wormhole and exact distance rankings in each CM/anatomical comparison. b, Agreement on the nearest match is a separate, less stable quantity. c, P7-to-P14 relative CM maturation-rank gaps for exact matches, Wormhole matches and all possible pairs; maturation was not used to choose matches. Selected regions show closer ranks, but this is not universal or a developmental clock. d, Eligible FB counts in the 96 sampled neighborhoods show why fine fibroblast dynamics remain uncertain. The updated descriptions do not retrain the encoder or establish ancestry.'),
('S4','results/cardiac_developmental_story/figures/Figure_2_regional_cardiomyocyte_programs','Measured CM regional patterns give anatomical comparisons to follow',
 'a,b, Native Slit2 expression at P7/P14, with reference endocardial cells shown in red. c,d, Native Ankrd1 expression at P7/P14, with operational arterial structures outlined. e, Released Slit2-CM and working-CM labels differ in distance to endocardial ECs; this is not independent anatomical validation. f, Earlier locally matched near-artery/control Ankrd1 RNA ratios at all four ages. These are partial spatial associations, not proof of vascular lineage or mechanical function. Kcnd2/Irx5 were not measured by the deposited Xenium panel. Coordinates and earlier estimates are retained.')]

def heading(doc,title):
 p=doc.add_heading(title,1);p.paragraph_format.page_break_before=True
 return p
def figure(doc,item):
 number,path,title,legend=item;p=doc.add_heading('Figure '+number+'. '+title,2);p.paragraph_format.page_break_before=True
 with Image.open(ROOT/(path+'.png')) as im:w,h=im.size
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4);p.add_run().add_picture(str(ROOT/(path+'.png')),width=Inches(min(7.0,7.05*w/h)))
 p=doc.add_paragraph(legend)
 for r in p.runs:r.font.size=Pt(9)
 p.paragraph_format.space_after=Pt(0)
def copy_evidence_tables(doc):
 # Only table/text sections are copied; report narratives and obsolete figures are not appended.
 src=Document(ROOT/'results/cardiac_fibroblast_CM_refit/Fibroblast_CM_spatial_reassessment.docx');active=False
 for node in src.element.body:
  text=''.join(node.itertext()) if node.tag.endswith('}p') else ''
  if 'What changed after the refit?' in text:active=True
  if active and 'Figure S1.' in text:break
  if active:doc.element.body.insert(len(doc.element.body)-1,copy.deepcopy(node))
def roadmap_table(doc):
 heading(doc,'Experiments that distinguish the developmental explanations')
 style.table(doc,['Question','Measurement','What it would resolve'],[
 ['Does vascular delivery precede metabolic change?','Independent intermediate-stage hearts; perfusion, local oxygen and CM energy use in matched ventricular regions','Blood-supply support versus shared regional development'],
 ['Is the P14 endocardial structural setting real?','Clear membranes, CM sarcomere RNA/protein, endocardial identity and comparable loading/anatomy','Within-CM structural variation versus transcript ownership or regional mechanics'],
 ['Does the P0 matrix candidate affect structure?','Eln/Fbln1/Fbln2 RNA and matrix protein, followed by controlled matrix/CM experiments','A regional marker pattern versus a matrix-related structural contribution'],
 ['Can early FB environments alter metabolism?','Same starting CM state, oxygen and endothelial support across defined FB/matrix conditions','A matrix contribution independent of vascular support'],
 ['Do early states contribute to later groups?','Denser spatial sampling, P3/P5/P10 hearts and validated lineage tracing','Related RNA states versus actual cell ancestry']])
 doc.add_paragraph('Functional CM measurements should include sarcomere organization, calcium handling, contraction and energy use. Completed division is a separate outcome from cell-cycle RNA. Exact transport and Wormhole nominate developmental comparisons; they do not replace these measurements.')
def main():
 OUT.mkdir(exist_ok=True,parents=True);text=SOURCE.read_text();title=text.splitlines()[0][2:];body,methods=text.split('## Methods',1)
 doc=Document();doc.core_properties.title=title;doc.core_properties.author='Heart atlas analysis';doc.core_properties.subject='Unified biological report integrating current CM, fibroblast, endothelial and cross-age spatial evidence'
 doc.core_properties.created=datetime.datetime(2026,10,27,12,tzinfo=datetime.timezone.utc);doc.core_properties.modified=doc.core_properties.created
 s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=Inches(.65);s.bottom_margin=Inches(.65);s.left_margin=Inches(.7);s.right_margin=Inches(.7);s.header_distance=Inches(.25);s.footer_distance=Inches(.25);s.different_first_page_header_footer=True
 n=doc.styles['Normal'];n.font.name='Cambria';n.font.size=Pt(11);n.paragraph_format.line_spacing=1.02;n.paragraph_format.space_after=Pt(4)
 for key in ['Title','Heading 1','Heading 2']:doc.styles[key].font.name='Arial';doc.styles[key].font.color.rgb=RGBColor.from_string('203C54')
 doc.styles['Heading 1'].font.size=Pt(14.5);doc.styles['Heading 2'].font.size=Pt(11.5)
 s.header.paragraphs[0].text='POSTNATAL HEART | INTEGRATED MUSCLE, MATRIX AND VASCULAR DEVELOPMENT'
 for r in s.header.paragraphs[0].runs:r.font.size=Pt(8)
 footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 doc.add_paragraph('UNIFIED BIOLOGICAL REPORT • CURRENT INTEGRATED RESULTS','Subtitle');doc.add_paragraph(title,'Title');doc.add_paragraph('P0–P21 mouse-heart snRNA-seq and Xenium analysis')
 doc.add_heading('Central message',2)
 doc.add_paragraph('The postnatal heart contains changing muscle, matrix and vascular environments. Early EC and CM cell-cycle activity occurs together, while later vascular proximity accompanies selected metabolic or structural RNA patterns. Fibroblast programs change and overlap, but the latest refit does not establish one general fibroblast effect on CM maturation. Specific structural settings now provide focused hypotheses to validate.')
 doc.add_paragraph('This single report integrates the CM states, five endothelial categories, alternative fibroblast identities, full-RNA matrix features, latest spatial refit, calibrated cross-age comparisons and experimental roadmap. Nine main figure plates, four supporting plates and consolidated tables are included. Explanations lead with biology; reproducible definitions are collected in Methods.')
 doc.add_paragraph('The newest fibroblast labels are incorporated into the CM-neighbor and matched fibroblast spatial analyses. Earlier capillary-adjacency estimates retain their original cohorts. Cross-age fibroblast descriptions are updated; exact/Wormhole matching plans remain frozen. One available section per age limits developmental and causal inference.')
 doc.add_paragraph('Prepared 27 October 2026 • Earlier reports and source data remain available.')
 heading(doc,'Abstract');add_blocks(doc,body.split('## Abstract',1)[1].lstrip())
 for item in FIGURES:figure(doc,item)
 add_blocks(doc,'## Methods'+methods)
 in_methods=False
 for p in doc.paragraphs:
  if p.style.name=='Heading 1':in_methods=p.text=='Methods'
  elif in_methods and p.style.name=='Normal':p.paragraph_format.space_after=Pt(4)
 heading(doc,'Cell identities and the biological questions they support')
 style.table(doc,['CM state','Interpretation to investigate','Important limit'],[
 ['Working ventricular','Contractile ventricular reference across stages','One label can contain different maturation-associated states'],
 ['Ankrd1-associated','Regional structural or loading-related adaptation','Not exclusively peri-arterial or proof of disease'],
 ['Slit2-associated','Chamber-facing regional specialization','Requires anatomical validation; Kcnd2/Irx5 unmeasured'],
 ['Cycling-associated','Early growth-related CM state','Cell-cycle RNA does not establish completed division'],
 ['Ddc-associated','Distinct measured RNA state','Specific tissue function remains unresolved']])
 style.table(doc,['EC subtype','Tissue role to examine','Important limit'],[
 ['Capillary-like','Coronary exchange network and muscle support','Proximity is not measured perfusion'],
 ['Endocardial-like','Chamber-lining environment and local structural relationships','Distinct from coronary capillaries'],
 ['Arterial-like','Supply-vessel environment and mural context','RNA label does not trace a vessel wall'],
 ['Venous-like','Venous vascular setting','Sparse calls do not establish absence'],
 ['Lymphatic-like','Lymphatic vascular setting','Requires multi-marker endothelial/lymphatic support']])
 doc.add_paragraph('The proposed roles guide comparisons. They are not new functional validation or a complete anatomical lineage map. Broad and fine assignments retain their source confidence and boundary-review fields.')
 copy_evidence_tables(doc)
 heading(doc,'Gfpt2 identity: evidence levels and reference checks')
 style.table(doc,['P21 evidence','Profiles or measurement','Meaning'],[
 ['Measured Gfpt2/Pi16/Adgrd1 pattern','1,339','Broader shared expression'],
 ['Also enriched for Col8a1 pattern','327','Overlap, not demonstrated transition'],
 ['Public Gfpt2 group with pattern flag','1,058 of 2,716','Mixed support within a public label'],
 ['Naming-gene-hidden majority','280','Grouping formed without the naming genes'],
 ['Hidden Gfpt2/Pi16 subsequently detected','About 84% / 91%','Consistent marker evidence in that selected subset'],
 ['Broader candidate / crosschecked subset','85 / 75','Different confidence layers, not total abundance']])
 doc.add_paragraph('Opposite-library tests show that recovering more Gfpt2 reference nuclei can also increase incorrect assignments of other families. A missing family in a same-age reference is not a biological zero. Candidate identity, a measured marker pattern and developmental fate remain separate questions.')
 heading(doc,'Which analyses are current in this combined report?')
 style.table(doc,['Analysis','Status'],[
 ['Broad cardiac/EC maps and CM subtype/maturity maps','Reviewed existing definitions retained'],
 ['Alternative fibroblast candidate/program maps','Current assignments and measured programs included'],
 ['CM–FB and endocardial/capillary fraction models','Latest refit with current FB labels/programs'],
 ['P14 matched Col8a1/Gfpt2 arterial comparison','Updated to broader/crosschecked labels'],
 ['Earlier matched capillary-adjacency and nuclear cycling comparisons','Earlier validated cohorts/estimates retained, not rerun'],
 ['Cross-age FB descriptions','Updated on the 96 frozen neighborhoods'],
 ['Exact/Wormhole matching plans and encoder','Frozen; no retraining or cell tracking'],
 ['This unified document','Presentation synthesis; no additional biological model fitted']])
 for item in SUPPLEMENT:figure(doc,item)
 heading(doc,'How to read the cell-boundary results')
 for para in SEGMENTATION_NOTE:doc.add_paragraph(para)
 doc.add_paragraph('The P14 structural association is stronger in core CM RNA than deposited whole-cell counts. The report retains that sensitivity rather than treating either view as a verified boundary. Block intervals do not integrate all segmentation uncertainty.')
 roadmap_table(doc)
 doc.save(OUT/'Postnatal_heart_unified_biological_report.docx')
 inventory=[];md=text+'\n\n## Figure plates\n\n'
 for number,path,title,legend in FIGURES+SUPPLEMENT:
  rel=os.path.relpath(ROOT/(path+'.png'),OUT);md+=f'### Figure {number}. {title}\n\n![Figure {number}]({rel})\n\n{legend}\n\n'
  inventory.append({'number':number,'source_stem':path,'title':title,'legend':legend})
 (OUT/'Report_text.md').write_text(md.rstrip()+'\n');(OUT/'figure_inventory.json').write_text(json.dumps(inventory,indent=2)+'\n')
 (OUT/'report_structure.json').write_text(json.dumps({'title':doc.core_properties.title,'main_text_words':len(body.split()),'methods_references_words':len(methods.split()),
  'main_figures':9,'supporting_figures':4,'latest_FB_refit_incorporated':True,'updated_arterial_comparison_incorporated':True,'transport_descriptors_updated':True,
  'earlier_capillary_comparisons_retained_with_original_cohorts':True,'new_biological_models_fitted_for_this_document':False,'new_biological_replicates':0},indent=2)+'\n')
 print('UNIFIED_WORD_CREATED',flush=True)
if __name__=='__main__':main()
