"""Write the focused biological report with integrated current figure plates."""
from pathlib import Path
import csv, json, datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from write_cardiac_biological_story import add_blocks
from write_cardiac_story import table

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'results/CM_maturation21_nonCM_context'
SOURCE=ROOT/'analysis/CM_maturation21_context_report.md'
STEM='CM_maturation_and_fibroblast_development'
FIGURES=[
 ('Figure_1_direct_21gene_CM_maturation_maps','CM maturation varies within the original cardiac sections',
  'a–d, Directly measured 21-gene core-RNA scores on one frozen reference scale across P0–P21; extreme colors are clipped at pooled 2nd/98th percentiles for display. e–h, Relative score ranks within the same age and CM identity. Only frozen strict ventricular map cohorts are colored. Percentile is a relative rank, not a percentage of physiological maturity. Gray retains tissue context, with display-only background subsampling. Counts and analyses use all eligible profiles. Scale bars, 500 µm.'),
 ('Figure_2_nonCM_subtypes_over_CM_maturation','Non-CM identities provide the context for CM developmental variation',
  'a–d, Five reviewed EC subtypes. e–h, Updated multi-gene fibroblast candidates, retaining Col8a1 as one family. Gfpt2/Pi16 names the existing Gfpt2-associated signature without demanding detection of both markers. i–l, Pericytes, smooth-muscle, immune and epicardial cells. Darker CM shading indicates higher measured 21-gene scores on the same scale across sections. Colored source points require broad source quality control. Unassigned profiles remain in background context rather than being forced into a subtype. The overlays do not establish direct contact or unbiased cell abundance. Scale bars, 500 µm.'),
 ('Figure_3_nonCM_relationships_with_CM_components','Neighbor relationships differ across developmental components',
  'a–d, Overall 21-gene, metabolic-regulatory, sarcomere and separate cell-cycle outcomes. Values are standardized within-region, CM-identity-conditioned associations from 50-µm neighborhoods. Asterisks indicate global BH q < 0.05 among all 235 supported primary tests, including the excitation–contraction outcome supplied in tables. Gray denotes unavailable or insufficiently supported comparisons. Broad source fractions use resolved-neighbor denominators; subtype fractions use all eligible members of their lineage, including unassigned profiles. Colors do not establish causal influence. One section is available at each age.'),
 ('Figure_4_capillary_score_and_robustness_tests','Capillary apposition does not establish uniformly greater CM maturity',
  'a, Overall score with the original capillary sources in core and deposited CM RNA, and newly matched QC-supported capillary sources in core RNA. b,c, Metabolic-regulatory and cell-cycle outcomes with original versus QC-supported sources. d, P7 gene-group omissions retain frozen normalization/reference; original-source pairs are unchanged. e, Equal-depth sensitivity compares the same depth-eligible pairs before and after exact subsampling to 50 molecules, averaged across three draws. f, Original-source sarcomere and excitation–contraction outcomes. Intervals are 95% spatial-block intervals, not animal replication. New QC-source cohorts contain 338, 514, 913 and 680 pairs. Source definitions change matched cohorts as well as capillary identities. The 20 original-source strict-ventricular core tests form the corrected primary matched family; sensitivity panels do not create additional discoveries.'),
 ('Figure_5_maturation_score_depth_and_gene_audit','Gene omission and capture sensitivity address different questions',
  'a,b, Mean age-related differences among identical depth-eligible CMs within each age, using original versus fixed-depth scores or jointly omitting Cebpb/Rxrg/Lpin1. Each curve is centered on its own P0 mean. c, Additive gene contributions to the original 21-gene P21-minus-P0 score difference in the full strict map cohorts; the three prespecified omitted genes are highlighted. d, Display-only samples of 1,500 CMs per stage show score versus captured transcript depth. e, Mean and range of correlations across 21 single-gene omissions, separately by age and RNA view. f, Change in buffered spatially held-out R² after adding fine identities or identities plus fibroblast programs to the anatomy/capture/broad-context model. High omission correlation does not establish functional validity, and subsampling discards biological information as well as equalizing molecule number.'),
 ('Figure_6_fibroblast_matrix_development','Fibroblast development changes matrix composition rather than uniformly increasing collagen RNA',
  'a–c, Corrected full-RNA library pseudobulk matrix scores within reviewed fibroblast families. Small points show the two libraries; lines show their unweighted means. Different gene sets have different expression baselines and are not protein measurements. d, Selected within-family gene changes; the displayed Fn1/Col1/Col3 directions agree across library cross-comparisons. Missing screened contrasts are not imputed as zero. e, Selected existing marker-withheld RNA similarities with reviewed family names. Dashed arrows indicate expression resemblance, not descendants, probabilities or conversion rates. f, Median void/outer-surface geometry within a conservative ventricular selector; only assessable groups with at least 30 profiles are shown. This is not validated transmural depth and does not track migration.')]

def rows(name):
    with open(OUT/name) as f:return list(csv.DictReader(f,delimiter='\t'))
def heading(doc,text):
    p=doc.add_heading(text,1);p.paragraph_format.page_break_before=True;return p
def main():
    source=SOURCE.read_text();body,methods=source.split('## Methods',1)
    title=source.splitlines()[0][2:];doc=Document();doc.core_properties.title=title
    doc.core_properties.author='Heart atlas analysis';doc.core_properties.subject='Measured CM maturation, non-CM neighborhoods and fibroblast developmental hypotheses'
    now=datetime.datetime(2026,10,27,16,tzinfo=datetime.timezone.utc)
    doc.core_properties.created=now;doc.core_properties.modified=now
    s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.left_margin=s.right_margin=Inches(.7);s.top_margin=s.bottom_margin=Inches(.65)
    s.header_distance=s.footer_distance=Inches(.25);s.different_first_page_header_footer=True
    n=doc.styles['Normal'];n.font.name='Cambria';n.font.size=Pt(11)
    n.paragraph_format.line_spacing=1.0;n.paragraph_format.space_after=Pt(4)
    for name in ['Title','Heading 1','Heading 2']:
        doc.styles[name].font.name='Arial';doc.styles[name].font.color.rgb=RGBColor.from_string('203C54')
    doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 2'].font.size=Pt(11)
    s.header.paragraphs[0].text='POSTNATAL HEART | MEASURED MATURATION AND THE CELLULAR ENVIRONMENT'
    for run in s.header.paragraphs[0].runs:run.font.size=Pt(8)
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    page=OxmlElement('w:fldSimple');page.set(qn('w:instr'),'PAGE');footer._p.append(page)
    doc.add_paragraph('FOCUSED BIOLOGICAL REASSESSMENT • 27 OCTOBER 2026','Subtitle')
    doc.add_paragraph(title,'Title')
    doc.add_paragraph('Direct Xenium 21-gene scores, updated non-CM maps, capillary-source sensitivity and answers to the fibroblast developmental questions')
    doc.add_heading('Direct answer',2)
    doc.add_paragraph('At P7, proximity to capillaries does not clearly establish greater overall cardiomyocyte maturity. The 21-gene comparison is small and uncertain, and the earlier lower-cycling relationship weakens when capillary sources are defined more strictly. The updated maps and component analyses identify specific developmental settings without claiming a general maturation effect from every neighboring subtype.')
    table(doc,['Stage','Mapped ventricular CMs','Original-source matched pairs','QC-source matched pairs'],[
        ['P0','17,272','577','338'],['P7','30,485','1,108','514'],['P14','33,659','1,250','913'],['P21','37,949','1,621','680']])
    doc.add_paragraph('Six integrated figure plates, biological interpretation, source tables and reproducible Methods are included. The earlier reports are preserved. One spatial section per age limits developmental and causal inference.')
    heading(doc,'Biological findings and interpretation')
    add_blocks(doc,body)
    heading(doc,'Selected quantitative comparisons')
    data=rows('matched_capillary_maturation_results.tsv');table_rows=[]
    for stage in ['P0','P7','P14','P21']:
        for cohort,label in [('strict_ventricular_pairs','Original capillaries'),('QC_capillary_strict_pairs','QC capillaries')]:
            r=next(r for r in data if r['stage']==stage and r['cohort']==cohort and r['outcome']=='core_Z__maturation21')
            table_rows.append([stage,label,r['matched_pairs'],f'{float(r["standardized_difference"]):+.3f}',f'{float(r["CI_low"]):+.3f} to {float(r["CI_high"]):+.3f}'])
    table(doc,['Stage','Capillary source','Pairs','21-gene difference (SD)','95% block interval'],table_rows)
    doc.add_paragraph('Apposed minus matched comparison. The original-source P7 overall score has q = 0.533 in the separate 20-test matched family. QC-source results are sensitivity comparisons with different matched cells; intervals excluding zero would not automatically establish corrected discoveries.')
    doc.add_heading('What do the fibroblast observations establish?',2)
    table(doc,['Question','Current answer'],[
        ['Keep Col8a1 as one family?','Yes. Svep1/Hmcn2 variation is not promoted to separate spatial families.'],
        ['Use Gfpt2/Pi16 as a family name?','Yes, as a multi-gene working identity; no obligatory double-positive gate.'],
        ['Does Postn decline establish cell depletion?','No. Marker repression, changing state, expansion of other populations and actual cell loss remain distinct possibilities.'],
        ['Does later position require migration?','No. Wall growth and remodeling can change distances. Validated equivalent anatomical coordinates are needed.'],
        ['Do all early fibroblasts become more collagen-rich?','No. Early cells already express collagen; several later populations reduce fibrillar-collagen RNA while changing matrix composition.'],
        ['Does the analysis establish Postn-to-Col8a1/Gfpt2 ancestry?','No. Existing RNA similarities nominate connections for temporal sampling and lineage tracing.'],
        ['Do finer identities explain the overall CM score better?','Not in the tested buffered spatial prediction models. Specific structural, metabolic or innervation outcomes may be more informative.']])
    for number,(stem,title,legend) in enumerate(FIGURES,1):
        p=doc.add_heading(f'Figure {number}. {title}',2);p.paragraph_format.page_break_before=True
        image=OUT/'figures'/(stem+'.png')
        with Image.open(image) as im:w,h=im.size
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4)
        p.add_run().add_picture(str(image),width=Inches(min(7.0,7.15*w/h)))
        p=doc.add_paragraph(legend);p.paragraph_format.space_after=Pt(0)
        for run in p.runs:run.font.size=Pt(9)
    add_blocks(doc,'## Methods'+methods)
    for t in doc.tables:
        for row in t.rows:
            prop=row._tr.get_or_add_trPr();prop.append(OxmlElement('w:cantSplit'))
    doc.save(OUT/(STEM+'.docx'))
    legends=[{'number':i+1,'stem':a,'title':b,'legend':c} for i,(a,b,c) in enumerate(FIGURES)]
    (OUT/'figure_legends.json').write_text(json.dumps(legends,indent=2)+'\n')
    md=source+'\n\n## Integrated figure plates\n\n'
    for r in legends:md+=f'### Figure {r["number"]}. {r["title"]}\n\n![Figure {r["number"]}](figures/{r["stem"]}.png)\n\n{r["legend"]}\n\n'
    (OUT/'Report_text.md').write_text(md.rstrip()+'\n')
    print('MATURATION21_WORD_CREATED',flush=True)

if __name__=='__main__':main()
