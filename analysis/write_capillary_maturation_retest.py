"""Portable Word/PDF documents from the same biological narrative and plate."""
from pathlib import Path
import re,html,zipfile,json,hashlib
from docx import Document
from docx.shared import Inches,Pt
from PIL import Image as PILImage
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak,Image
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import fitz

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'results/capillary_maturation_retest'
def main():
 text=(OUT/'Report_text.md').read_text();title=text.splitlines()[0][2:]
 doc=Document();sec=doc.sections[0];sec.left_margin=sec.right_margin=Inches(.65)
 sec.top_margin=sec.bottom_margin=Inches(.6)
 doc.styles['Normal'].font.name='Cambria';doc.styles['Normal'].font.size=Pt(11)
 doc.styles['Normal'].paragraph_format.space_after=Pt(6)
 doc.core_properties.title=title
 for name,file in [('D','DejaVuSans.ttf'),('DB','DejaVuSans-Bold.ttf')]:pdfmetrics.registerFont(TTFont(name,'/usr/share/fonts/truetype/dejavu/'+file))
 styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='BodyD',fontName='D',fontSize=10,leading=13.4,spaceAfter=8))
 styles.add(ParagraphStyle(name='TitleD',fontName='DB',fontSize=18,leading=23,spaceAfter=12))
 styles.add(ParagraphStyle(name='HeadD',fontName='DB',fontSize=13,leading=17,spaceBefore=10,spaceAfter=7,keepWithNext=True))
 story=[]
 def pp(s,style='BodyD'):return Paragraph(html.escape(s),styles[style])
 for block in re.split(r'\n\s*\n',text.strip()):
  if block.startswith('# '):doc.add_heading(block[2:],0);story.append(pp(block[2:],'TitleD'))
  elif block.startswith('## '):
   if block=='## Methods and reproducibility':
    doc.add_page_break();story.append(PageBreak())
    doc.add_heading('Integrated figure',1);doc.add_picture(str(OUT/'Capillary_maturation_integrated_figure.png'),width=Inches(7))
    legend=(OUT/'figure_legend.txt').read_text();doc.add_paragraph(legend)
    with PILImage.open(OUT/'Capillary_maturation_integrated_figure.png') as im:w,h=im.size
    width=min(490,490*w/h)
    story.extend([pp('Integrated figure','HeadD'),Image(str(OUT/'Capillary_maturation_integrated_figure.png'),width=width,height=width*h/w),pp(legend)])
    doc.add_page_break();story.append(PageBreak())
   doc.add_heading(block[3:],1);story.append(pp(block[3:],'HeadD'))
  elif block.startswith('|'):
   rows=[[x.strip() for x in line.strip('|').split('|')] for line in block.splitlines() if not re.match(r'^\|[- :|]+\|$',line)]
   tb=doc.add_table(rows=0,cols=len(rows[0]));tb.style='Light Shading Accent 1'
   for row in rows:
    cells=tb.add_row().cells
    for cell,val in zip(cells,row):cell.text=val
   pdf=Table([[pp(x) for x in row] for row in rows],colWidths=[45,105,340],repeatRows=1)
   pdf.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e6edf2')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOX',(0,0),(-1,-1),.3,colors.gray),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]));story.extend([pdf,Spacer(1,10)])
  else:
   s=' '.join(block.splitlines());doc.add_paragraph(s);story.append(pp(s))
 doc.save(OUT/'Capillary_maturation_retest.docx')
 def footer(canvas,document):canvas.setFont('D',8);canvas.drawString(45,20,'POSTNATAL HEART | CAPILLARY RETEST');canvas.drawRightString(565,20,str(document.page))
 SimpleDocTemplate(str(OUT/'Capillary_maturation_retest.pdf'),pagesize=(612,792),rightMargin=45,leftMargin=45,topMargin=38,bottomMargin=38).build(story,onFirstPage=footer,onLaterPages=footer)
 with zipfile.ZipFile(OUT/'Capillary_maturation_retest.docx') as z:
  assert z.testzip() is None
  media=[z.read(n) for n in z.namelist() if n.startswith('word/media/')];assert len(media)==1
  assert hashlib.sha256(media[0]).hexdigest()==hashlib.sha256((OUT/'Capillary_maturation_integrated_figure.png').read_bytes()).hexdigest()
 pdf=fitz.open(OUT/'Capillary_maturation_retest.pdf');figurepages=[]
 for i,page in enumerate(pdf):
  assert len(page.get_text().strip())>100
  if page.get_image_info():figurepages.append(i+1)
  for b in page.get_text('dict')['blocks']:
   if b['type']!=0:continue
   for line in b['lines']:
    for span in line['spans']:
     r=fitz.Rect(span['bbox']);assert r.x0>=35 and r.y0>=10 and r.x1<=580 and r.y1<=780,(i,span['text'],r)
 assert len(figurepages)==1
 (OUT/'document_validation.json').write_text(json.dumps(dict(status='PASS',PDF_pages=len(pdf),figure_pages=figurepages,Word_image_hash_matches=True,PDF_text_bounds_checked=True),indent=2)+'\n')
if __name__=='__main__':main()
