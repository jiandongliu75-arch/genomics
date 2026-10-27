"""Combine the uploaded reports in clear biological prose with current FB data."""
from pathlib import Path
import csv,json,os,datetime
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
import write_cardiac_story as style
from write_cardiac_biological_story import add_blocks

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'results/heart_atlas_clear_language_revision'
SOURCE=ROOT/'analysis/heart_atlas_clear_language_revision.md'
BIO='results/cardiac_biological_synthesis/figures/';FB='results/fibroblast_multimethod_identity/figures/'
FIGURES=[
('1',BIO+'Figure_1_CM_identity_and_maturation','CMs of the same subtype show different maturation-associated RNA patterns',
 'a,b, Ventricular-associated CM labels at P0 and P7. c,d, Relative position within the section’s 21-gene maturation-score distribution. e,f, Score distributions within each CM subtype. A high score indicates stronger expression of the measured panel, rather than a direct measurement of CM function. Labels and coordinates are retained from the earlier analysis. Scale bars, 500 µm.'),
('2','results/heart_atlas_clear_language_revision/figures/Figure_2_fibroblast_matrix_patterns','Fibroblasts develop different matrix-associated features',
 'a, RNA programs across ages and fibroblast groups, averaged over the two libraries and shown relative to the other displayed groups. b, Gene patterns in named P21 RNA subgroups; dot size shows detection and color shows relative expression. c, Earlier P14 matched Col8a1/Gfpt2 comparison: proximity to arterial structures and the smooth-muscle share of neighboring cells are separate measures. These values use the earlier conservative labels, not the new classifier labels. d, Selected RNA-signature similarities with interpreted markers withheld. Dashed arrows mark comparisons to test; they do not trace cell descendants.'),
('3',FB+'Figure_3_stage_maps_and_boundary_stability','Alternative methods recover broader fibroblast candidate maps',
 'a–d, Candidate fibroblast identities on the four original sections. e, Assignment coverage among eligible broad fibroblasts. Coverage is not subtype abundance or an independently measured accuracy. f, P21 agreement between whole-cell and boundary-core labels, compared with random loss of whole-cell transcripts to similar depth. This helps separate loss of information from sensitivity to transcript location. Nuclear alternatives remain review flags. Gray retains profiles without a confident candidate; scale bars, 500 µm.'),
('4',FB+'Figure_2_P21_identity_and_program_maps','A Gfpt2-associated pattern extends beyond the narrow subtype',
 'a, Public named fibroblast labels within the eligible P21 cohort. b, Multi-method cell candidates. c, The subset also supported in the cell core and after naming markers are hidden. d, Direct Gfpt2/Pi16/Adgrd1 marker score relative to matched control genes. e,f, Gfpt2/Col8a1 pattern combinations across eligible fibroblasts or within the public Gfpt2 subset. A pattern-positive profile is not automatically a member of one subtype. Public valve labels are not equated with Kirrel3 candidates. Gray and review flags retain uncertainty; scale bars, 500 µm.'),
('5',BIO+'Figure_4_vascular_growth_and_CM_specialization','Vascular growth and CM maturation have different relationships over time',
 'a, Regional vascular sequence from Lu et al. (2021), rather than a new measurement of blood flow. b, Earlier P7 cycling-CM/Postn-fibroblast metabolic association under different source rules and distances. This result becomes less convincing when the full set of tested relationships is considered. c, Earlier P14 Ankrd1-CM/endocardial-EC sarcomere association. Intervals summarize uncertainty within a section. d, Earlier matched comparison of CMs alongside capillary ECs with comparable CMs without that proximity; this uses a broader CM cohort. e, Vessel growth, local access and functional support are different measurements. The newest classifier-based FB labels have not been used to refit these estimates.'),
('6',BIO+'Figure_5_CM_identity_conditioned_transport','Some cellular environments resemble one another across ages',
 'a, Comparison design, restricted to the same CM subtype and a similar position relative to the endocardium. b, Earlier strict fibroblast/endothelial descriptors in 96 observed neighborhoods. c,d, Agreement between Wormhole and exact neighborhood distances. e, Similarity of environments around the same or different CM subtypes within an age. f, Relative maturation-rank agreement was checked after matching, rather than used to choose matches. Results vary by region and age. New FB labels do not alter these plans; no individual cells or descendants were followed.'),
('7',BIO+'Figure_6_regional_developmental_model','Experiments can separate blood supply, matrix and contact effects',
 'a, Proposed sequence of early growth followed by later specialization. The arrows show developmental windows, not observed cell conversion. b, Blood supply, matrix/mechanics and cellular contact are different possible contributions to CM development. c, Independent intermediate-stage hearts, controlled functional experiments and lineage tracing can distinguish them. Direct measurements should include perfusion, CM energy use, calcium handling and contraction. This is a model to test, rather than a completed causal experiment.')]
SUPPLEMENT=[
('S1',FB+'Figure_1_reference_recovery_and_sensitivity','The reference tests show why more labels need not mean better assignments',
 'a, P21 Gfpt2 precision and recall in opposite-library RNA tests; each point is one held-out library. b, Mean RNA-label recovery by family and method. c, Sensitivity to hidden naming markers and pooled late references. d, Sensitivity to reduced RNA information. These tests recover our reviewed RNA definitions; they do not establish independent Xenium accuracy. Some methods recover more Gfpt2 nuclei while also assigning other families incorrectly.'),
('S2',FB+'Figure_4_heldout_genes_and_published_programs','Hidden genes help check assignments, while marker patterns can overlap',
 'a, Measured P21 expression after grouping with classifiers that could not use the seven naming/anchor genes. Printed values show detection percentages. b, Gfpt2/Col8a1 pattern combinations among eligible profiles with public Gfpt2 labels. c, Published marker coverage in the shared panel. The valve set has too few measured genes for a multi-gene module. d, A candidate subtype, a measured marker pattern and developmental fate describe different kinds of evidence.')]

SEGMENTATION_NOTE=[
 'Cell boundaries determine which transcripts are assigned to each spatial profile. The atlas combines membrane, cell-interior and nuclear stains to help identify these boundaries. This is more appropriate for large and multinucleated cardiomyocytes than simply drawing a fixed region around each nucleus.',
 'The authors compared automatic boundaries with expert outlines in ten fields: five at P0 and five at P21. Multimodal segmentation performed better than nucleus expansion. However, its boundary-agreement score was only about 45–50% at P0 and around 10% at P21 under a demanding 90% overlap rule. These are approximate readings from the published plot. A low object score does not mean the same proportion of transcripts were assigned incorrectly: an outline can fail the strict overlap test while retaining many correctly assigned interior transcripts.',
 'The poorer P21 agreement matters when interpreting small fibroblasts or endothelial cells beside transcript-rich cardiomyocytes. At the same time, a CM profile without a visible nucleus can be a valid section through a large cell. Removing all such profiles would discard real tissue information.',
 'Our approach keeps whole-cell, inner-cell and nuclear evidence visible. Nuclear agreement can strengthen a label. A lack of informative nuclear transcripts is not sufficient to reject it, while a supported alternative nuclear identity is a reason for review. None of these views is treated as anatomical ground truth.',
 'The next useful validation is to examine selected interfaces directly, with clear membrane outlines and several subtype markers in the same cells. This can test whether an apparent shared RNA pattern lies within one fibroblast or across a boundary. The full earlier segmentation assessment and the original paper’s Methods/Supplementary Figure 1 remain linked in the source record.'
]

def readrows(path):
    with open(ROOT/path) as f:return list(csv.DictReader(f,delimiter='\t'))
def figure(doc,item):
    number,path,title,legend=item;p=doc.add_heading(f'Figure {number}. {title}',2);p.paragraph_format.page_break_before=True
    picture=ROOT/(path+'.png')
    with Image.open(picture) as im:w,h=im.size
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4)
    p.add_run().add_picture(str(picture),width=Inches(min(7.0,7.10*w/h)))
    p=doc.add_paragraph(legend);p.paragraph_format.space_after=Pt(0)
    for r in p.runs:r.font.size=Pt(9)
def table_page(doc,title):
    p=doc.add_heading(title,1);p.paragraph_format.page_break_before=True
def main():
    OUT.mkdir(parents=True,exist_ok=True);text=SOURCE.read_text();title=text.splitlines()[0][2:]
    doc=Document();doc.core_properties.title=title;doc.core_properties.author='Heart atlas analysis'
    doc.core_properties.subject='Clear biological synthesis incorporating alternative fibroblast identity evidence'
    doc.core_properties.created=datetime.datetime(2026,10,26,12,tzinfo=datetime.timezone.utc);doc.core_properties.modified=doc.core_properties.created
    s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=Inches(.65);s.bottom_margin=Inches(.65);s.left_margin=Inches(.7);s.right_margin=Inches(.7)
    s.header_distance=Inches(.25);s.footer_distance=Inches(.25);s.different_first_page_header_footer=True
    n=doc.styles['Normal'];n.font.name='Cambria';n.font.size=Pt(11);n.paragraph_format.line_spacing=1.04;n.paragraph_format.space_after=Pt(5)
    for name in ['Title','Heading 1','Heading 2']:
        doc.styles[name].font.name='Arial';doc.styles[name].font.color.rgb=RGBColor.from_string('203C54')
    doc.styles['Heading 1'].font.size=Pt(14.5);doc.styles['Heading 2'].font.size=Pt(11.5)
    s.header.paragraphs[0].text='POSTNATAL HEART DEVELOPMENT | MUSCLE, MATRIX AND BLOOD VESSELS'
    for r in s.header.paragraphs[0].runs:r.font.size=Pt(8)
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.add_paragraph('INTEGRATED BIOLOGICAL REPORT • CLEAR-LANGUAGE REVISION','Subtitle');doc.add_paragraph(title,'Title')
    doc.add_paragraph('P0–P21 mouse-heart snRNA-seq and Xenium analysis')
    doc.add_heading('Central message',2)
    doc.add_paragraph('The findings suggest that muscle, matrix and blood vessels develop together, but different aspects of cardiomyocyte maturation emerge at different times and in different regions. Fibroblast subtype identities and the gene patterns shared between them provide complementary ways to examine this process.')
    doc.add_paragraph('This report combines the biological synthesis with the alternative fibroblast identity reassessment. Seven main figure plates and two supporting plates are included. The explanations lead with the biology; detailed analysis rules are collected in Methods.')
    doc.add_paragraph('The new fibroblast maps have not yet been used to recalculate the earlier CM-neighbor associations or cross-age transport results. One available section per age limits developmental and causal conclusions.')
    doc.add_paragraph('Prepared 26 October 2026 • Source results and earlier reports are retained for traceability.')
    p=doc.add_heading('Abstract',1);p.paragraph_format.page_break_before=True
    body,methods=text.split('## Methods',1)
    # The first heading is already placed with an explicit page break.
    body=body.split('## Abstract',1)[1].lstrip();add_blocks(doc,body)
    for item in FIGURES:figure(doc,item)
    add_blocks(doc,'## Methods'+methods)
    table_page(doc,'What the evidence supports')
    style.table(doc,['Biological observation','What it suggests','What remains to test'],[
        ['CMs of one subtype have different maturation RNA patterns','Subtype and developmental state are different features','CM energy use, calcium handling and contraction'],
        ['Later FB groups differ in matrix-associated RNA','Different contributions to tissue organization are plausible','Matrix protein, deposition and mechanical properties'],
        ['Gfpt2 pattern extends beyond the narrow family','Shared programs can cross subtype boundaries','Membrane-resolved same-cell marker combinations'],
        ['Muscle and endothelial cell-cycle RNA vary together early','Coordinated local growth','Actual completed division and temporal order'],
        ['Capillary proximity relates to different CM components at different ages','Growth and later vascular support may accompany different steps','Perfusion, oxygen delivery and contact effects'],
        ['P14 Ankrd1 CMs near endocardial-like ECs show more sarcomere RNA','A subtype-specific structural relationship','Regional anatomy, loading and endothelial signals'],
        ['Some early/late RNA states or neighborhoods resemble each other','Developmental comparisons worth following','Independent intermediate stages and ancestry tracing']])
    table_page(doc,'Fibroblast identities and programs: reading the numbers')
    coverage=readrows('results/fibroblast_multimethod_identity/final_candidate_coverage.tsv')
    rows=[]
    for family in ['Postn-associated','Gfpt2-associated','Col8a1-associated','Lsamp-associated','Kirrel3-associated']:
        row=[family.replace('-associated','')]
        for stage in ['P0','P7','P14','P21']:
            r=next(r for r in coverage if r['family']==family and r['stage']==stage);row.append(f'{int(r["candidate_profiles"]):,}')
        rows.append(row)
    style.table(doc,['Candidate family','P0','P7','P14','P21'],rows)
    doc.add_paragraph('Counts refer to candidate assignments among eligible broad fibroblasts. A zero means no candidates under that reference and rule; it is not proof that every cell or marker of that kind is absent. These are not tissue-wide abundance estimates. Kirrel3 candidates are not assigned valve anatomy from their name.')
    style.table(doc,['P21 Gfpt2 evidence','Profiles','Meaning'],[
        ['Marker-pattern enrichment','1,339','Broader measured Gfpt2/Pi16/Adgrd1 signal under the module rule'],
        ['Also enriched for Col8a1 pattern','327','Overlapping RNA patterns; not proof of transition'],
        ['Public Gfpt2 profiles with Gfpt2-pattern flag','1,058 of 2,716','The public-labelled population contains varied marker evidence'],
        ['Naming-marker-hidden majority label','280','Candidate grouping made without the seven naming/anchor genes'],
        ['Cell-supported candidate','85','Reference-screened model agreement without a supported core conflict'],
        ['Cell/core and naming-marker consistency','75','A narrower supported subset; not the total biological population']])
    doc.add_heading('Which results include the new fibroblast labels?',2)
    style.table(doc,['Result','Status in this revision'],[
        ['Fibroblast identity and marker-pattern maps','Updated alternative classifiers and direct programs included'],
        ['CM-neighbor association estimates','Earlier models retained; not refitted with the newest classifiers'],
        ['P14 arterial-proximity comparison','Earlier conservative matched subset retained'],
        ['Cross-age neighborhood matching','Earlier representation and plans retained']])
    for item in SUPPLEMENT:figure(doc,item)
    table_page(doc,'How to read the cell-boundary results')
    for para in SEGMENTATION_NOTE:doc.add_paragraph(para)
    doc.add_paragraph('Sources: Wang et al. (2026), segmentation Methods and Supplementary Figure 1; the earlier detailed segmentation assessment accompanies the source record.')
    name='Postnatal_heart_integrated_clear_language_report';doc.save(OUT/(name+'.docx'))
    inventory=[]
    md=text+'\n\n## Figure plates\n\n'
    for number,path,title,legend in FIGURES+SUPPLEMENT:
        rel=os.path.relpath(ROOT/(path+'.png'),OUT)
        md+=f'### Figure {number}. {title}\n\n![Figure {number}]({rel})\n\n{legend}\n\n'
        inventory.append({'number':number,'source_stem':path,'title':title,'legend':legend,'new_plot':path.startswith('results/heart_atlas_clear_language_revision/')})
    md+='## Reading the cell-boundary results\n\n'+'\n\n'.join(SEGMENTATION_NOTE)+'\n'
    (OUT/'Report_text.md').write_text(md.rstrip()+'\n')
    (OUT/'figure_inventory.json').write_text(json.dumps(inventory,indent=2)+'\n')
    (OUT/'report_structure.json').write_text(json.dumps({'title':doc.core_properties.title,'main_figures':7,'supporting_figures':2,'main_text_words':len(body.split()),
      'new_analysis_models_fitted':False,'new_classifier_labels_used_to_refit_CM_associations':False,'uploaded_reports_integrated':True},indent=2)+'\n')
    print('CLEAR_INTEGRATED_REPORT_CREATED',OUT/(name+'.docx'),flush=True)
if __name__=='__main__':main()
