"""Create matched Word and PDF biological reports from one narrative and plates."""
from pathlib import Path
import csv,json,re,html,datetime,hashlib,zipfile
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak,Image,Table,TableStyle,KeepTogether
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image as PILImage
from write_cardiac_biological_story import add_blocks
from write_cardiac_story import table

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'results/joint_RNA_Xenium_signaling'
SOURCE=ROOT/'analysis/joint_signaling_report.md';STEM='Joint_RNA_Xenium_CM_development'
def rows(name):
 with open(OUT/name) as f:return list(csv.DictReader(f,delimiter='\t'))
def inline(text):
 text=html.escape(text)
 text=re.sub(r'\*\*(.+?)\*\*',r'<b>\1</b>',text)
 text=re.sub(r'`(.+?)`',r'\1',text)
 return text
def main():
 source=SOURCE.read_text();title=source.splitlines()[0][2:]
 narrative,methods=source.split('## Methods and interpretation limits',1)
 legends=json.loads((OUT/'figure_legends.json').read_text())
 cohorts=rows('analysis_cohorts.tsv');primary=rows('incremental_prediction_gains.tsv')
 chosen=[r for r in primary if r['primary_analysis']=='True' and r['comparison']=='E minus C_fine_context' and r['outcome']=='core_Z__maturation21']
 quantitative=[[r['stage'],f"{float(r['delta_R2']):+.4f}",f"{float(r['CI_low']):+.4f} to {float(r['CI_high']):+.4f}",f"{float(r['primary_grid_BH_q']):.3g}"] for r in chosen]
 doc=Document();doc.core_properties.title=title;doc.core_properties.author='Heart atlas analysis'
 doc.core_properties.created=doc.core_properties.modified=datetime.datetime(2026,10,27,20,tzinfo=datetime.timezone.utc)
 sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.left_margin=sec.right_margin=Inches(.7);sec.top_margin=sec.bottom_margin=Inches(.65)
 sec.header_distance=sec.footer_distance=Inches(.25)
 normal=doc.styles['Normal'];normal.font.name='Cambria';normal.font.size=Pt(11)
 normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1
 for name in ['Title','Heading 1','Heading 2']:
  doc.styles[name].font.name='Arial';doc.styles[name].font.color.rgb=RGBColor.from_string('203C54')
 doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 2'].font.size=Pt(11)
 sec.header.paragraphs[0].text='POSTNATAL HEART | JOINT RNA AND SPATIAL ANALYSIS'
 for run in sec.header.paragraphs[0].runs:run.font.size=Pt(8)
 page=OxmlElement('w:fldSimple');page.set(qn('w:instr'),'PAGE');sec.footer.paragraphs[0]._p.append(page)
 doc.add_paragraph('NEW ANALYSIS • 27 OCTOBER 2026','Subtitle');doc.add_paragraph(title,'Title')
 add_blocks(doc,'\n'.join(narrative.splitlines()[1:]))
 p=doc.add_heading('What was tested',1)
 table(doc,['Age','Mapped ventricular CMs','CMs retained for spatial tests','All cardiac source profiles'],
  [[r['stage'],f"{int(r['mapped_CMs']):,}",f"{int(r['test_eligible_CMs']):,}",f"{int(r['all_cardiac_source_profiles']):,}"] for r in cohorts if r['view']=='core'])
 doc.add_paragraph('The same source profiles and receiver CMs are used for core and deposited-RNA comparisons. Boundary buffers exclude CMs from model fitting; they do not remove those cells from the original atlas. One observed spatial section is available at each age.')
 table(doc,['Age','Overall 21-gene ΔR²','95% spatial-block interval','Primary BH q'],quantitative)
 doc.add_paragraph('Non-CM expression plus collective transport (E) compared with receptor, CM-to-CM and fine-context controls (C_fine_context). Positive values indicate improved prediction, not a positive effect on maturation. Intervals are conditional on fitted models within one section.')
 for i,entry in enumerate(legends,1):
  p=doc.add_heading(f'Figure {i}. {entry["title"]}',2);p.paragraph_format.page_break_before=True
  image=OUT/'figures'/(entry['stem']+'.png')
  with PILImage.open(image) as im:w,h=im.size
  doc.add_picture(str(image),width=Inches(min(7.0,7.1*w/h)))
  p=doc.add_paragraph(entry['legend'])
  for run in p.runs:run.font.size=Pt(9)
 p=doc.add_heading('Methods and interpretation limits',1);p.paragraph_format.page_break_before=True
 add_blocks(doc,methods)
 for t in doc.tables:
  for row in t.rows:row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 doc.save(OUT/(STEM+'.docx'))

 # The PDF is laid out from the identical narrative, tables and image files.
 font='/usr/share/fonts/truetype/dejavu/'
 for name,file in [('D','DejaVuSans.ttf'),('DB','DejaVuSans-Bold.ttf'),('DI','DejaVuSans-Oblique.ttf'),('DBI','DejaVuSans-BoldOblique.ttf')]:pdfmetrics.registerFont(TTFont(name,font+file))
 pdfmetrics.registerFontFamily('D',normal='D',bold='DB',italic='DI',boldItalic='DBI')
 styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='BodyD',fontName='D',fontSize=10.2,leading=13.4,spaceAfter=7))
 styles.add(ParagraphStyle(name='TitleD',fontName='DB',fontSize=19,leading=24,textColor=colors.HexColor('#203C54'),spaceAfter=14))
 styles.add(ParagraphStyle(name='HeadD',fontName='DB',fontSize=13,leading=17,spaceBefore=10,spaceAfter=8,textColor=colors.HexColor('#203C54'),keepWithNext=True))
 styles.add(ParagraphStyle(name='SubD',fontName='DB',fontSize=10.5,leading=14,spaceBefore=8,spaceAfter=5,keepWithNext=True))
 styles.add(ParagraphStyle(name='CaptionD',fontName='D',fontSize=8.3,leading=10.8,spaceAfter=6))
 styles.add(ParagraphStyle(name='TableD',fontName='D',fontSize=8.4,leading=10.7))
 story=[]
 def para(text,style='BodyD'):return Paragraph(inline(text),styles[style])
 def blocks(text):
  for block in re.split(r'\n\s*\n',text.strip()):
   if not block.strip():continue
   if block.startswith('## '):story.append(para(block[3:],'HeadD'))
   elif block.startswith('### '):story.append(para(block[4:],'SubD'))
   elif block.startswith('- '):
    for line in block.splitlines():story.append(para('• '+line.lstrip('- ')))
   else:story.append(para(' '.join(block.splitlines())))
 def pdf_table(headers,data,widths):
  cells=[[para(str(x),'TableD') for x in row] for row in [headers]+data]
  t=Table(cells,colWidths=widths,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([
   ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e6edf2')),('VALIGN',(0,0),(-1,-1),'TOP'),
   ('GRID',(0,0),(-1,-1),.3,colors.HexColor('#c7d2dc')),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),
   ('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]));story.append(t);story.append(Spacer(1,10))
 story.extend([para('NEW ANALYSIS • 27 OCTOBER 2026','CaptionD'),para(title,'TitleD')]);blocks('\n'.join(narrative.splitlines()[1:]))
 story.append(para('What was tested','HeadD'))
 pdf_table(['Age','Mapped ventricular CMs','CMs retained for spatial tests','All cardiac source profiles'],
  [[r['stage'],f"{int(r['mapped_CMs']):,}",f"{int(r['test_eligible_CMs']):,}",f"{int(r['all_cardiac_source_profiles']):,}"] for r in cohorts if r['view']=='core'],[48,142,157,160])
 story.append(para('The same source profiles and receiver CMs are used for core and deposited-RNA comparisons. Boundary buffers exclude CMs from model fitting; they do not remove those cells from the original atlas. One observed spatial section is available at each age.'))
 pdf_table(['Age','Overall 21-gene ΔR²','95% spatial-block interval','Primary BH q'],quantitative,[48,142,207,110])
 story.append(para('Non-CM expression plus collective transport (E) compared with receptor, CM-to-CM and fine-context controls (C_fine_context). Positive values indicate improved prediction, not a positive effect on maturation. Intervals are conditional on fitted models within one section.'))
 for i,entry in enumerate(legends,1):
  story.extend([PageBreak(),para(f'Figure {i}. {entry["title"]}','HeadD')]);image=OUT/'figures'/(entry['stem']+'.png')
  with PILImage.open(image) as im:w,h=im.size
  width=min(507,510*w/h);story.extend([Image(str(image),width=width,height=width*h/w),Spacer(1,7),para(entry['legend'],'CaptionD')])
 story.extend([PageBreak(),para('Methods and interpretation limits','HeadD')]);blocks(methods)
 def footer(canvas,doc):
  canvas.saveState();canvas.setFont('D',7.5);canvas.setFillColor(colors.HexColor('#526575'))
  canvas.drawString(50,772,'POSTNATAL HEART | JOINT RNA AND SPATIAL ANALYSIS');canvas.drawRightString(560,23,str(doc.page));canvas.restoreState()
 SimpleDocTemplate(str(OUT/(STEM+'.pdf')),pagesize=(612,792),leftMargin=50,rightMargin=55,topMargin=44,bottomMargin=42,
  title=title,author='Heart atlas analysis').build(story,onFirstPage=footer,onLaterPages=footer)
 md=source+'\n\n## Figure plates\n\n'
 for i,r in enumerate(legends,1):md+=f'### Figure {i}. {r["title"]}\n\n![Figure {i}](figures/{r["stem"]}.png)\n\n{r["legend"]}\n\n'
 (OUT/'Report_text.md').write_text(md.rstrip()+'\n')
 print('JOINT_WORD_PDF_CREATED',flush=True)
if __name__=='__main__':main()
