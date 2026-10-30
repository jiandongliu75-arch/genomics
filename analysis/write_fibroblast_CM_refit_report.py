"""Write the completed reassessment in observation-led biological prose."""
from pathlib import Path
import json,csv,datetime,os
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
import write_cardiac_story as style
from write_cardiac_biological_story import add_blocks
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'results/cardiac_fibroblast_CM_refit';SOURCE=ROOT/'analysis/fibroblast_CM_refit_report.md'
LEGENDS=[
('1','Figure_1_fibroblast_states_and_shared_programs','Fibroblast identities and shared patterns give complementary maps',
 'a–d, Broader candidate fibroblast identities on the original P0–P21 sections. Gray retains unassigned fibroblasts and tissue context; background cells may be subsampled for display only. e, Candidate counts, with unavailable same-age reference classes marked N/A. This does not mean the corresponding genes are absent. Counts measure assignment coverage, not accuracy or tissue abundance. f, Measured patterns can be scored independently of family availability. Values summarize each section relative to matched control genes; they do not measure matrix protein or function. Scale bars, 500 µm.'),
('2','Figure_2_measured_fibroblast_program_maps','Early and late fibroblast patterns overlap without requiring a subtype label',
 'a–d, Postn/Ptn/Tnc/Col24a1 pattern. e–h, Gfpt2/Pi16/Adgrd1 pattern. i–l, Col8a1/Cmklr1/Col6a3/Htra3 pattern. Each row uses a common color scale across ages, showing eligible fibroblasts regardless of their candidate subtype. Negative scores mean expression ranks below matched controls, not negative RNA or biological absence. One pattern can occur across several subtype labels. Background tissue is shown in gray. Scale bars, 500 µm.'),
('3','Figure_3_local_program_relationships','Local matrix relationships differ by CM state and outcome',
 'a–c, Within-region associations between neighboring fibroblast patterns and CM energy-metabolism, sarcomere or cell-cycle expression. Working, Ankrd1 and cycling-associated CMs are displayed; all eligible CM states enter the source tables and correction. Red and blue indicate positive and negative estimates. Dots mark intervals excluding zero before multiple-testing correction; none of these fibroblast-program associations passes the global primary correction. Gray indicates an unavailable or insufficiently supported comparison. The models account for vascular surroundings and capture. Each age has one section.'),
('4','Figure_4_early_fibroblast_CM_relationships','Early fibroblast environments suggest hypotheses without establishing metabolic instruction',
 'a, The P7 Postn-identity/metabolic relationship in cycling-associated CMs weakens with broader and crosschecked assignments. b,c, Postn- and Gfpt2-associated patterns in whole-cell, inner-cell, deposited and depth-matched fibroblast RNA, with an additional CM-gene-excluded scoring check. Negative and positive estimates suggest contrasting environments but do not establish opposing functions. d, The 50-/30-µm comparison uses the same 2,753 cycling-associated CMs. Smaller neighborhoods do not provide clear corresponding support. Points and intervals describe within-section estimates and 95% spatial-block intervals; they are not independent-heart replication.'),
('5','Figure_5_structural_relationships_and_boundaries','Structural RNA relationships depend on CM identity and transcript ownership',
 'a, At P14, endocardial-like surroundings relate to higher sarcomere expression in Ankrd1 CMs, with a smaller relationship in working CMs. Removing Myl2 retains the direction; the overall score shows no corresponding increase. b, The relationship remains after adjustment for fibroblast programs but weakens in deposited CM RNA. The paired 50-/30-µm comparison uses the same 4,528 Ankrd1 CMs. c, A P0 elastic-matrix/sarcomere candidate retains its direction across fibroblast RNA views and CM-gene-excluded scoring, but does not pass the primary correction. d, The P21 basement-matrix estimate weakens with inner-cell fibroblast RNA. Intervals describe uncertainty within sections, conditional on the chosen count views.'),
('6','Figure_6_spatial_context_and_cross_age_comparisons','Fibroblast position and changing surroundings guide specific developmental questions',
 'a, P14 Col8a1/Gfpt2 arterial proximity after the same technical matching design, using earlier stringent, broader and crosschecked identities. The contrast becomes smaller with broader assignments. Matching permits reused controls. b, Arterial proximity among candidate families; N/A marks unavailable families. Kirrel3 candidates are not assigned valve anatomy. c, Updated fibroblast patterns on frozen exact P7-to-P14 neighborhood matches, grouped by CM identity and endocardial distance. Only matching mass with informative fibroblasts at both ends contributes; supported mass is supplied. d, Eligible fibroblasts among the sampled sources: the median is two, and 15 of 96 neighborhoods have none. These comparisons do not trace cell transitions.'),
('S1','Figure_S1_fibroblast_identity_associations','Subtype labels do not establish one uniform effect on neighboring CMs',
 'a–c, Candidate-family associations with CM metabolic, sarcomere and cell-cycle expression. The 50-µm models compare CMs within regions and account for vascular surroundings and capture. Gray marks unavailable or insufficiently supported comparisons. No fibroblast identity association passes the global primary correction. The full source tables also include other eligible CM states, the 21-gene and excitation–contraction outcomes, alternative label layers, smaller neighborhoods and transcript-ownership sensitivities.')]
def rows(name):
 with open(OUT/name) as f:return list(csv.DictReader(f,delimiter='\t'))
def heading(doc,title):
 p=doc.add_heading(title,1);p.paragraph_format.page_break_before=True
 return p
def figure(doc,item):
 number,stem,title,legend=item;p=doc.add_heading('Figure '+number+'. '+title,2);p.paragraph_format.page_break_before=True
 picture=OUT/'figures'/(stem+'.png')
 with Image.open(picture) as im:w,h=im.size
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4)
 p.add_run().add_picture(str(picture),width=Inches(min(7.0,7.15*w/h)))
 p=doc.add_paragraph(legend)
 for r in p.runs:r.font.size=Pt(9)
 p.paragraph_format.space_after=Pt(0)
def main():
 text=SOURCE.read_text();body,methods=text.split('## Methods',1);title=text.splitlines()[0][2:]
 doc=Document();doc.core_properties.title=title;doc.core_properties.author='Heart atlas analysis'
 doc.core_properties.subject='Updated fibroblast identities, measured programs and cardiomyocyte relationships'
 doc.core_properties.created=datetime.datetime(2026,10,26,16,tzinfo=datetime.timezone.utc);doc.core_properties.modified=doc.core_properties.created
 s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=Inches(.65);s.bottom_margin=Inches(.65);s.left_margin=Inches(.7);s.right_margin=Inches(.7);s.header_distance=Inches(.25);s.footer_distance=Inches(.25);s.different_first_page_header_footer=True
 n=doc.styles['Normal'];n.font.name='Cambria';n.font.size=Pt(11);n.paragraph_format.line_spacing=1.02;n.paragraph_format.space_after=Pt(4)
 for name in ['Title','Heading 1','Heading 2']:
  doc.styles[name].font.name='Arial';doc.styles[name].font.color.rgb=RGBColor.from_string('203C54')
 doc.styles['Heading 1'].font.size=Pt(14.5);doc.styles['Heading 2'].font.size=Pt(11.5)
 s.header.paragraphs[0].text='POSTNATAL HEART | FIBROBLAST ENVIRONMENTS AND MUSCLE DEVELOPMENT'
 for r in s.header.paragraphs[0].runs:r.font.size=Pt(8)
 footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
 doc.add_paragraph('COMPLETED SPATIAL REASSESSMENT • BIOLOGICAL REPORT','Subtitle');doc.add_paragraph(title,'Title')
 doc.add_paragraph('P0–P21 mouse-heart snRNA-seq and Xenium analysis')
 doc.add_heading('Central message',2)
 doc.add_paragraph('Fibroblast patterns change as the heart develops, and a Col8a1-associated component remains more often near arteries. These differences do not establish one general fibroblast effect on CM maturation. The refit instead identifies specific structural and metabolic settings to test, with the strongest corrected association concerning P14 Ankrd1 CMs near the endocardial environment.')
 doc.add_paragraph('The revised fibroblast assignments are now included in newly fitted CM-neighbor models. Whole-cell, inner-cell and deposited RNA, smaller neighborhoods and shared-boundary comparisons distinguish stronger observations from tentative candidates.')
 doc.add_paragraph('Six main figure plates, one supporting plate, source tables and detailed Methods are included. Cross-age neighborhood descriptions are updated; the original exact-transport and Wormhole matches are retained. One section per age limits developmental and causal conclusions.')
 doc.add_paragraph('Prepared 26 October 2026 • Earlier reports and original data are preserved.')
 heading(doc,'Abstract');add_blocks(doc,body.split('## Abstract',1)[1].lstrip())
 for item in LEGENDS[:-1]:figure(doc,item)
 add_blocks(doc,'## Methods'+methods)
 in_methods=False
 for p in doc.paragraphs:
  if p.style.name=='Heading 1':in_methods=p.text=='Methods'
  elif in_methods and p.style.name=='Normal':p.paragraph_format.space_after=Pt(4)
 heading(doc,'What changed after the refit?')
 style.table(doc,['Question','Updated observation','Biological interpretation'],[
 ['Does the P7 Postn relationship persist?','Negative metabolic estimate becomes smaller and uncertain with broader labels; smaller-neighborhood support is weak','A developing-matrix hypothesis remains; suppression by Postn fibroblasts is not established'],
 ['Do late Col8a1 fibroblasts favor arteries?','About 11% near arteries versus 7% of matched Gfpt2 candidates; direction retained in the crosschecked subset','A vessel-associated component within a broader matrix population'],
 ['Does one FB program predict general CM maturation?','No FB identity/program association passes the full primary correction','Different components and CM states require separate tests'],
 ['Which FB candidate deserves focused imaging?','P0 elastic-matrix expression relates to lower working-CM sarcomere RNA across several sensitivity checks','Test matrix organization and CM structure; mechanism remains unknown'],
 ['What is the clearest corrected relationship?','P14 endocardial-like surroundings and Ankrd1-CM sarcomere RNA; weaker in deposited CM counts','Validate anatomy and transcript ownership before testing contact or loading'],
 ['What does the temporal comparison add?','P7-to-P14 matched surroundings lose Postn-associated and gain Col8a1-associated expression','A developmental comparison to sample more densely, not a tracked conversion']])
 heading(doc,'Fibroblast assignments and directly measured patterns')
 cov=rows('fibroblast_assignment_coverage.tsv');table=[]
 for fam in ['Postn','Gfpt2','Col8a1','Lsamp','Kirrel3']:
  row=[fam]
  for stage in ['P0','P7','P14','P21']:
   r=next(r for r in cov if r['stage']==stage and r['family']==fam and r['label_layer']=='candidate');row.append(f'{int(float(r["candidate_profiles"])):,}' if r['reference_assessable']=='True' else 'Not assessable')
  table.append(row)
 style.table(doc,['Candidate family','P0','P7','P14','P21'],table)
 doc.add_paragraph('Not assessable means that this family is absent from the same-age reference. It is not a biological zero. Counts show assignment coverage among eligible fibroblasts, rather than tissue-wide abundance. Broader coverage is not independently established accuracy; apparent reference precision uses thresholds selected from the same reference predictions.')
 pc=rows('fibroblast_program_coverage.tsv');table=[]
 for mod in ['Postn.Fib','Gfpt2.Fib','Col8a1.Fib','Lsamp.Fib','Fibrillar_matrix','Basement_matrix','Elastic_matrix']:
  row=[mod.replace('.Fib',' pattern').replace('_',' ')]
  for stage in ['P0','P7','P14','P21']:
   r=next(r for r in pc if r['stage']==stage and r['program']==mod and r['view']=='refined');row.append(f'{int(r["enriched_profiles"]):,}')
  table.append(row)
 style.table(doc,['Pattern-enriched profiles','P0','P7','P14','P21'],table)
 doc.add_paragraph('These descriptive flags require the measured multi-gene pattern to exceed matched controls and at least two markers to be detected. They are not subtype counts, functional activity or evidence of biological absence when low. Different patterns can occur in the same cell. Eligible broad FB denominators are 2,020, 7,365, 8,098 and 10,867.')
 heading(doc,'Selected estimates and exact program definitions')
 r=rows('CM_fibroblast_refitted_coefficients.tsv');table=[]
 selected=[('P7','p.CM','candidate__Postn','core_Z__metabolic','Postn identity → cycling-CM metabolic RNA'),('P7','p.CM','refined__Gfpt2.Fib','core_Z__metabolic','Gfpt2 pattern → cycling-CM metabolic RNA'),('P0','v.CM','refined__Elastic_matrix','core_Z__sarcomere','Elastic matrix → working-CM sarcomere RNA'),('P21','Ankrd1.CM','refined__Basement_matrix','core_Z__sarcomere','Basement matrix → Ankrd1 sarcomere RNA'),('P14','Ankrd1.CM','EC__Endocardial-like','core_Z__sarcomere','Endocardial context → Ankrd1 sarcomere RNA')]
 for stage,state,target,outcome,label in selected:
  a=next(x for x in r if x['stage']==stage and x['CM_state']==state and x['target']==target and x['outcome']==outcome and x['primary_analysis']=='True')
  table.append([stage+' · '+label,f'{float(a["coefficient"]):+.3f}',f'{float(a["CI_low"]):+.3f} to {float(a["CI_high"]):+.3f}',f'{float(a["primary_grid_BH_q"]):.3g}'])
 style.table(doc,['Primary comparison','Estimate','95% block interval','Global BH q'],table)
 doc.add_paragraph('Values are within-stage standardized associations, not measured functional effects. Intervals describe spatial variation conditional on the selected masks and model. The global correction covers all 925 primary tests. Sensitivity estimates are supplied separately and do not create extra corrected discoveries.')
 defs=rows('fibroblast_program_definitions.tsv');style.table(doc,['Measured pattern','Genes'],[[r['program'].replace('.Fib',' pattern').replace('_',' '),r['genes'].replace(';',', ')] for r in defs])
 doc.add_paragraph('Fibroblast genes do not overlap the CM maturation panel. Individual program models do not establish effects independent of the other patterns; joint-program sensitivities are supplied. Marker-pattern expression is not a measurement of matrix deposition or mechanics.')
 figure(doc,LEGENDS[-1])
 heading(doc,'Experiments suggested by the reassessment')
 style.table(doc,['Priority','Measurement','What it would distinguish'],[
 ['1. Validate the P14 structural setting','Clear membranes; Ankrd1 and sarcomere RNA/protein; endocardial identity; matched inner-wall anatomy','True within-CM structural variation versus transcript assignment or regional loading'],
 ['2. Examine the P0 elastic-matrix candidate','Eln/Fbln1/Fbln2 RNA and matrix protein; CM sarcomere organization in comparable regions','A shared regional feature versus a local matrix-related structural relationship'],
 ['3. Test early metabolic environments','Compare fibroblast/matrix conditions with the same starting CM state, oxygen and endothelial support','Matrix contribution independent of vascular support'],
 ['4. Resolve developmental order','Independent P3/P5/P10 hearts with perfusion, oxygen, matrix and CM function measurements','Whether improved blood delivery, matrix organization or contact precedes each CM change'],
 ['5. Follow cell states through time','Denser spatial neighborhoods and validated lineage tracing','Related expression states versus actual cell ancestry']])
 doc.add_paragraph('One section per age cannot establish population-level developmental timing. Shared-boundary and core-RNA analyses do not prove functional contact, oxygen delivery or complete removal of mixed RNA. The experiments should measure contraction, calcium handling, energy use and completed division separately from RNA scores.')
 doc.save(OUT/'Fibroblast_CM_spatial_reassessment.docx')
 md=text+'\n\n## Figure plates\n\n'
 for number,stem,title,legend in LEGENDS:md+=f'### Figure {number}. {title}\n\n![Figure {number}](figures/{stem}.png)\n\n{legend}\n\n'
 (OUT/'Report_text.md').write_text(md.rstrip()+'\n')
 (OUT/'figure_legends.json').write_text(json.dumps([{'number':a,'stem':b,'title':c,'legend':d} for a,b,c,d in LEGENDS],indent=2)+'\n')
 (OUT/'report_structure.json').write_text(json.dumps({'title':doc.core_properties.title,'main_text_words':len(body.split()),'methods_references_words':len(methods.split()),'main_figures':6,'supporting_figures':1,
  'new_CM_FB_models_fitted':True,'latest_FB_candidates_used':True,'exact_and_Wormhole_plans_retrained':False,'transport_descriptors_updated':True,'new_biological_replicates':0},indent=2)+'\n')
 print('REFIT_WORD_CREATED',flush=True)
if __name__=='__main__':main()
