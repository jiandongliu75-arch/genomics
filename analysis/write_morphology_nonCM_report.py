"""Matched Word and PDF from one narrative, table and two figure plates."""
from pathlib import Path
import re,json,html,hashlib,zipfile
from docx import Document
from docx.shared import Inches,Pt
from PIL import Image as PILImage
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak,Image,Table,TableStyle
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import fitz
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'results/morphology_stratified_nonCM';STEM='Morphology_nonCM_CM_development'
FOOTER_LABEL='POSTNATAL HEART | ANATOMY AND NON-CM CONTEXT'
METHODS_HEADING='## Methods and interpretation limits'
def main():
 text=(OUT/'Report_text.md').read_text();legends=json.loads((OUT/'figure_legends.json').read_text());doc=Document()
 sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.left_margin=sec.right_margin=Inches(.65);sec.top_margin=sec.bottom_margin=Inches(.6)
 doc.styles['Normal'].font.name='Cambria';doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.space_after=Pt(6)
 doc.core_properties.title=text.splitlines()[0][2:]
 for name,file in [('D','DejaVuSans.ttf'),('DB','DejaVuSans-Bold.ttf')]:pdfmetrics.registerFont(TTFont(name,'/usr/share/fonts/truetype/dejavu/'+file))
 styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='BodyD',fontName='D',fontSize=10,leading=13.5,spaceAfter=8))
 styles.add(ParagraphStyle(name='TitleD',fontName='DB',fontSize=18,leading=23,spaceAfter=12));styles.add(ParagraphStyle(name='HeadD',fontName='DB',fontSize=13,leading=17,spaceBefore=10,spaceAfter=7,keepWithNext=True))
 styles.add(ParagraphStyle(name='CaptionD',fontName='D',fontSize=8.5,leading=11,spaceAfter=7))
 story=[];expected=[]
 def pp(s,style='BodyD'):return Paragraph(html.escape(s),styles[style])
 for block in re.split(r'\n\s*\n',text.strip()):
  if block.startswith('# '):doc.add_heading(block[2:],0);story.append(pp(block[2:],'TitleD'))
  elif block.startswith('## '):
   if block==METHODS_HEADING:
    for i,entry in enumerate(legends,1):
     doc.add_page_break();story.append(PageBreak());heading=f'Figure {i}. {entry["title"]}';doc.add_heading(heading,1);story.append(pp(heading,'HeadD'))
     image=OUT/'figures'/(entry['stem']+'.png');expected.append(hashlib.sha256(image.read_bytes()).hexdigest())
     with PILImage.open(image) as im:w,h=im.size
     width=min(7.1,7.1*w/h);doc.add_picture(str(image),width=Inches(width));doc.add_paragraph(entry['legend'])
     width=min(510,540*w/h);story.append(Image(str(image),width=width,height=width*h/w));story.append(pp(entry['legend'],'CaptionD'))
    doc.add_page_break();story.append(PageBreak())
   doc.add_heading(block[3:],1);story.append(pp(block[3:],'HeadD'))
  elif block.startswith('|'):
   rows=[[v.strip() for v in line.strip('|').split('|')] for line in block.splitlines() if not re.match(r'^\|[- :|]+\|$',line)]
   tb=doc.add_table(rows=0,cols=len(rows[0]));tb.style='Light Shading Accent 1'
   for row in rows:
    cells=tb.add_row().cells
    for cell,value in zip(cells,row):cell.text=value
   widths={3:[50,215,215],4:[45,145,145,145],5:[40,130,130,100,80]}[len(rows[0])]
   pdf=Table([[pp(x) for x in row] for row in rows],colWidths=widths,repeatRows=1)
   pdf.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e6edf2')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOX',(0,0),(-1,-1),.3,colors.gray),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]));story.extend([pdf,Spacer(1,10)])
  else:
   s=' '.join(block.splitlines());doc.add_paragraph(s);story.append(pp(s))
 doc.save(OUT/(STEM+'.docx'))
 def footer(canvas,document):canvas.setFont('D',8);canvas.drawString(45,20,FOOTER_LABEL);canvas.drawRightString(565,20,str(document.page))
 SimpleDocTemplate(str(OUT/(STEM+'.pdf')),pagesize=(612,792),rightMargin=45,leftMargin=45,topMargin=38,bottomMargin=38).build(story,onFirstPage=footer,onLaterPages=footer)
 with zipfile.ZipFile(OUT/(STEM+'.docx')) as z:
  assert z.testzip() is None
  media={hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')};assert media==set(expected)
 pdf=fitz.open(OUT/(STEM+'.pdf'));figurepages=[]
 for i,page in enumerate(pdf):
  assert len(page.get_text().strip())>100
  if page.get_image_info():figurepages.append(i+1)
  for block in page.get_text('dict')['blocks']:
   if block['type']!=0:continue
   for line in block['lines']:
    for span in line['spans']:
     r=fitz.Rect(span['bbox']);assert r.x0>=35 and r.y0>=10 and r.x1<=580 and r.y1<=780,(i,span['text'],r)
 assert len(figurepages)==len(legends)
 (OUT/'document_validation.json').write_text(json.dumps(dict(status='PASS',PDF_pages=len(pdf),figure_pages=figurepages,Word_current_image_hashes_verified=True,PDF_text_bounds_checked=True),indent=2)+'\n')
if __name__=='__main__':main()
